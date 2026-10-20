import asyncio
import os
import re
import random
import hashlib
from pathlib import Path

import pymupdf as fitz

from docx import Document
from docx.shared import Pt

from aiogram import Bot, Dispatcher, F
from aiogram.filters import Command
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import (
    Message,
    CallbackQuery,
    InlineKeyboardMarkup,
    InlineKeyboardButton,
    FSInputFile,
)
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode


# ============================================================
# USTOZ AI
# ============================================================

TOKEN = os.getenv("BOT_TOKEN")

if not TOKEN:
    raise RuntimeError(
        "BOT_TOKEN topilmadi!\n"
        "PowerShell'da BOT_TOKEN ni o'rnating."
    )

BOOKS_DIR = Path("books")
BOOKS_DIR.mkdir(exist_ok=True)

bot = Bot(
    token=TOKEN,
    default=DefaultBotProperties(parse_mode=ParseMode.HTML)
)

dp = Dispatcher()


# ============================================================
# FSM
# ============================================================

class TeacherStates(StatesGroup):
    choosing_class = State()
    choosing_subject = State()
    choosing_book = State()
    entering_pages = State()


# ============================================================
# XOTIRA
# ============================================================

# Har bir foydalanuvchi uchun oxirgi yaratilgan testlar
USER_DATA = {}

# Takror savollarni kamaytirish uchun
USED_QUESTIONS = {}


# ============================================================
# MATN TOZALASH
# ============================================================

def clean_text(text: str) -> str:
    if not text:
        return ""

    text = text.replace("\u00a0", " ")
    text = text.replace("\r", "\n")

    # Juda ko'p bo'shliqlarni tozalash
    text = re.sub(r"[ \t]+", " ", text)

    # Ko'p bo'sh qatorlarni kamaytirish
    text = re.sub(r"\n\s*\n+", "\n", text)

    return text.strip()


def normalize_sentence(text: str) -> str:
    text = clean_text(text)

    # Sahifa raqami kabi juda qisqa narsalarni olib tashlash
    text = re.sub(r"^\d+$", "", text).strip()

    return text


# ============================================================
# SINF KLAVIATURASI
# ============================================================

def class_keyboard():
    buttons = []

    for i in range(1, 12):
        buttons.append(
            InlineKeyboardButton(
                text=f"🎓 {i}-sinf",
                callback_data=f"class:{i}"
            )
        )

    rows = []

    for i in range(0, len(buttons), 2):
        rows.append(buttons[i:i + 2])

    return InlineKeyboardMarkup(inline_keyboard=rows)


# ============================================================
# SUBJECT TOPISH
# ============================================================

def get_subjects(grade: int):
    grade_dir = BOOKS_DIR / f"{grade}-sinf"

    if not grade_dir.exists():
        return []

    subjects = []

    # Subfolderlar
    for item in sorted(grade_dir.iterdir()):
        if item.is_dir():
            pdfs = list(item.glob("*.pdf"))

            if pdfs:
                subjects.append(item.name)

    # To'g'ridan-to'g'ri PDF
    for pdf in sorted(grade_dir.glob("*.pdf")):
        subjects.append(pdf.stem)

    return subjects


# ============================================================
# BOOK TOPISH
# ============================================================

def get_books(grade: int, subject: str):
    grade_dir = BOOKS_DIR / f"{grade}-sinf"

    subject_dir = grade_dir / subject

    if subject_dir.exists() and subject_dir.is_dir():
        return sorted(subject_dir.glob("*.pdf"))

    direct_pdf = grade_dir / f"{subject}.pdf"

    if direct_pdf.exists():
        return [direct_pdf]

    # Xavfsizroq qidiruv
    result = []

    for pdf in grade_dir.glob("*.pdf"):
        if pdf.stem == subject:
            result.append(pdf)

    return result


# ============================================================
# SAHIFA FORMATINI TEKSHIRISH
# ============================================================

def parse_pages(value: str):
    value = value.strip()

    # 10-20
    match = re.fullmatch(r"(\d+)\s*-\s*(\d+)", value)

    if match:
        start = int(match.group(1))
        end = int(match.group(2))

        if start < 1 or end < 1:
            return None

        if start > end:
            return None

        if end - start + 1 > 100:
            return None

        return start, end

    # 10
    if value.isdigit():
        page = int(value)

        if page < 1:
            return None

        return page, page

    return None


# ============================================================
# PDF O'QISH
# ============================================================

def read_pdf(pdf_path: Path, start_page: int, end_page: int):
    try:
        document = fitz.open(pdf_path)

        total_pages = len(document)

        if total_pages == 0:
            document.close()
            return "", 0, "PDF bo'sh."

        if start_page > total_pages:
            document.close()
            return "", total_pages, (
                f"PDF faqat {total_pages} sahifadan iborat."
            )

        real_end = min(end_page, total_pages)

        all_text = []

        for page_number in range(start_page - 1, real_end):
            page = document[page_number]

            text = page.get_text("text")

            if text:
                all_text.append(text)

        document.close()

        result = clean_text("\n".join(all_text))

        return result, total_pages, None

    except Exception as e:
        return "", 0, f"PDF o'qishda xatolik: {e}"


# ============================================================
# GAP BO'LIB AJRATISH
# ============================================================

def split_into_sentences(text: str):
    text = clean_text(text)

    if not text:
        return []

    # Nuqta, ?, !
    parts = re.split(
        r"(?<=[.!?])\s+|\n+",
        text
    )

    sentences = []

    for part in parts:
        part = normalize_sentence(part)

        if not part:
            continue

        # Juda qisqa gaplar kerak emas
        if len(part) < 35:
            continue

        # Juda uzun gaplarni ham haddan tashqari ishlatmaymiz
        if len(part) > 500:
            part = part[:500]

        sentences.append(part)

    return sentences


# ============================================================
# SO'ZLAR
# ============================================================

STOP_WORDS = {
    "va",
    "ham",
    "bu",
    "shu",
    "bir",
    "bilan",
    "uchun",
    "ning",
    "ni",
    "ga",
    "da",
    "dan",
    "dir",
    "esa",
    "yoki",
    "ammo",
    "lekin",
    "ular",
    "u",
    "o",
    "men",
    "sen",
    "biz",
    "siz",
    "bo'ladi",
    "bo'lgan",
    "bo'lib",
    "edi",
    "ekan",
    "emas",
    "juda",
    "eng",
    "har",
    "har bir",
    "qanday",
    "qaysi",
    "barcha",
    "ko'p",
    "kam",
}


def extract_words(text: str):
    words = re.findall(
        r"[A-Za-zА-Яа-яЁёЎўҚқҒғҲҳ0-9]+(?:['’ʻ`-][A-Za-zА-Яа-яЁёЎўҚқҒғҲҳ0-9]+)*",
        text
    )

    result = []

    for word in words:
        clean = word.strip()

        if len(clean) < 3:
            continue

        if clean.lower() in STOP_WORDS:
            continue

        result.append(clean)

    return result


# ============================================================
# MUHIM SO'Z TANLASH
# ============================================================

def choose_answer_word(sentence: str):
    words = extract_words(sentence)

    if not words:
        return None

    # Uzunroq, mazmunliroq so'zlarni afzal ko'ramiz
    candidates = [
        word
        for word in words
        if 4 <= len(word) <= 25
    ]

    if not candidates:
        candidates = words

    # Eng uzun so'zlardan birini tanlash
    candidates.sort(
        key=lambda x: len(x),
        reverse=True
    )

    # Juda tasodifiy bo'lib ketmasligi uchun
    top = candidates[:min(5, len(candidates))]

    return random.choice(top)


# ============================================================
# DISTRAKTORLAR
# ============================================================

def make_distractors(
    correct_answer: str,
    source_text: str
):
    words = extract_words(source_text)

    unique = []

    correct_lower = correct_answer.lower()

    for word in words:
        if word.lower() == correct_lower:
            continue

        if word.lower() in {
            x.lower() for x in unique
        }:
            continue

        # Juda qisqa yoki juda uzun variantlarni kamaytiramiz
        if len(word) < 3 or len(word) > 30:
            continue

        unique.append(word)

    # O'xshash uzunlikdagi so'zlarni afzal ko'ramiz
    similar = [
        word
        for word in unique
        if abs(len(word) - len(correct_answer)) <= 6
    ]

    if len(similar) >= 3:
        unique = similar

    random.shuffle(unique)

    distractors = []

    for word in unique:
        if word.lower() == correct_lower:
            continue

        if word.lower() in {
            x.lower() for x in distractors
        }:
            continue

        distractors.append(word)

        if len(distractors) == 3:
            break

    return distractors


# ============================================================
# SAVOL HASH
# ============================================================

def question_hash(question: str, answer: str):
    raw = (
        question.strip().lower()
        + "|"
        + answer.strip().lower()
    )

    return hashlib.sha256(
        raw.encode("utf-8")
    ).hexdigest()


# ============================================================
# 20 TA TEST GENERATOR
# ============================================================

def generate_questions(
    source_text: str,
    count: int = 20,
    user_id: int = 0
):
    sentences = split_into_sentences(source_text)

    if len(sentences) < 5:
        return []

    used = USED_QUESTIONS.setdefault(user_id, set())

    # Har safar boshqa tartib
    random.shuffle(sentences)

    questions = []

    for sentence in sentences:

        if len(questions) >= count:
            break

        answer = choose_answer_word(sentence)

        if not answer:
            continue

        distractors = make_distractors(
            answer,
            source_text
        )

        if len(distractors) < 3:
            continue

        # Javob so'zini gap ichidan almashtirish
        pattern = re.escape(answer)

        question_text = re.sub(
            pattern,
            "_____",
            sentence,
            count=1
        )

        # Almashtirilmagan bo'lsa
        if question_text == sentence:
            continue

        # Savol hash
        q_hash = question_hash(
            question_text,
            answer
        )

        if q_hash in used:
            continue

        options = [
            answer,
            distractors[0],
            distractors[1],
            distractors[2],
        ]

        random.shuffle(options)

        correct_letter = "ABCD"[
            options.index(answer)
        ]

        question = {
            "question": question_text,
            "options": {
                "A": options[0],
                "B": options[1],
                "C": options[2],
                "D": options[3],
            },
            "answer": correct_letter,
            "answer_text": answer,
        }

        questions.append(question)
        used.add(q_hash)

    return questions


# ============================================================
# TESTNI MATNGA AYLANTIRISH
# ============================================================

def question_to_text(index: int, question: dict):
    result = []

    result.append(
        f"<b>{index}. {question['question']}</b>"
    )

    result.append(
        f"A) {question['options']['A']}"
    )

    result.append(
        f"B) {question['options']['B']}"
    )

    result.append(
        f"C) {question['options']['C']}"
    )

    result.append(
        f"D) {question['options']['D']}"
    )

    result.append(
        f"✅ <b>To'g'ri javob: "
        f"{question['answer']}) "
        f"{question['answer_text']}</b>"
    )

    return "\n".join(result)


# ============================================================
# UZUN XABARNI BO'LISH
# ============================================================

async def send_questions(
    message: Message,
    questions: list
):
    # 5 tadan yuboramiz
    # Telegram limitiga tushmasligi uchun

    for start in range(0, len(questions), 5):

        chunk = questions[start:start + 5]

        text = []

        for local_index, question in enumerate(
            chunk,
            start=start + 1
        ):
            text.append(
                question_to_text(
                    local_index,
                    question
                )
            )

        await message.answer(
            "\n\n".join(text)
        )


# ============================================================
# TEST UCHUN TUGMALAR
# ============================================================

def test_menu_keyboard():
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(
                    text="🔄 Yana 20 ta",
                    callback_data="again"
                )
            ],
            [
                InlineKeyboardButton(
                    text="📄 Word fayl",
                    callback_data="word"
                )
            ],
            [
                InlineKeyboardButton(
                    text="📚 Boshqa fan",
                    callback_data="subjects"
                )
            ],
            [
                InlineKeyboardButton(
                    text="🎓 Boshqa sinf",
                    callback_data="classes"
                )
            ],
            [
                InlineKeyboardButton(
                    text="🏠 Asosiy menyu",
                    callback_data="home"
                )
            ],
        ]
    )


# ============================================================
# WORD YARATISH
# ============================================================

def create_word_file(
    user_id: int,
    questions: list,
    grade: int,
    subject: str,
    book_name: str,
    page_range: str
):
    document = Document()

    # Sarlavha
    title = document.add_heading(
        "USTOZ AI — TEST",
        level=0
    )

    title.alignment = 1

    p = document.add_paragraph()

    p.add_run(
        f"Sinf: {grade}-sinf\n"
        f"Fan: {subject}\n"
        f"Darslik: {book_name}\n"
        f"Sahifa: {page_range}\n"
        f"Testlar soni: {len(questions)}"
    )

    document.add_paragraph("")

    for index, question in enumerate(
        questions,
        start=1
    ):
        p = document.add_paragraph()

        run = p.add_run(
            f"{index}. {question['question']}"
        )

        run.bold = True

        for letter in ["A", "B", "C", "D"]:
            document.add_paragraph(
                f"{letter}) {question['options'][letter]}"
            )

        answer = document.add_paragraph()

        answer_run = answer.add_run(
            f"To'g'ri javob: "
            f"{question['answer']}) "
            f"{question['answer_text']}"
        )

        answer_run.bold = True

        document.add_paragraph("")

    # Javoblar kaliti
    document.add_heading(
        "Javoblar kaliti",
        level=1
    )

    answer_text = []

    for index, question in enumerate(
        questions,
        start=1
    ):
        answer_text.append(
            f"{index}. {question['answer']}"
        )

    document.add_paragraph(
        " | ".join(answer_text)
    )

    # Shriftsiz qolmasin
    styles = document.styles

    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)

    filename = (
        f"ustoz_ai_{user_id}.docx"
    )

    filepath = Path(filename)

    document.save(filepath)

    return filepath


# ============================================================
# START
# ============================================================

@dp.message(Command("start"))
async def start_command(
    message: Message,
    state: FSMContext
):
    await state.clear()

    await message.answer(
        "🎓 <b>USTOZ AI</b>\n\n"
        "O'qituvchilar uchun aqlli test generator.\n\n"
        "Men darslikning siz ko'rsatgan sahifalarini "
        "o'qib, shu mavzudan test tayyorlayman.\n\n"
        "📚 Avval sinfni tanlang:",
        reply_markup=class_keyboard()
    )

    await state.set_state(
        TeacherStates.choosing_class
    )


# ============================================================
# RESTART
# ============================================================

@dp.message(Command("restart"))
async def restart_command(
    message: Message,
    state: FSMContext
):
    await state.clear()

    await message.answer(
        "🔄 <b>Jarayon qayta boshlandi!</b>\n\n"
        "🎓 Sinfni tanlang:",
        reply_markup=class_keyboard()
    )

    await state.set_state(
        TeacherStates.choosing_class
    )


# ============================================================
# HELP
# ============================================================

@dp.message(Command("help"))
async def help_command(message: Message):
    await message.answer(
        "📖 <b>USTOZ AI yordam</b>\n\n"
        "1️⃣ /start — botni boshlash\n"
        "2️⃣ Sinfni tanlash\n"
        "3️⃣ Fanni tanlash\n"
        "4️⃣ Darslikni tanlash\n"
        "5️⃣ Sahifa oralig'ini yozish\n"
        "   Masalan: <code>10-20</code>\n"
        "6️⃣ Bot 20 ta test tayyorlaydi\n"
        "7️⃣ Testni Word formatida olish mumkin\n\n"
        "🔄 /restart — jarayonni boshidan boshlash\n"
        "❓ /help — yordam"
    )


# ============================================================
# JOIN
# ============================================================

@dp.message(Command("join"))
async def join_command(message: Message):
    await message.answer(
        "🤖 <b>USTOZ AI</b>\n\n"
        "Bu bot o'qituvchilarga darslik asosida "
        "test tayyorlashda yordam beradi.\n\n"
        "📚 Darslik → Sahifa → Test → Word"
    )


# ============================================================
# CLASS TANLASH
# ============================================================

@dp.callback_query(
    TeacherStates.choosing_class,
    F.data.startswith("class:")
)
async def class_selected(
    callback: CallbackQuery,
    state: FSMContext
):
    grade = int(
        callback.data.split(":")[1]
    )

    subjects = get_subjects(grade)

    await state.update_data(
        grade=grade,
        subjects=subjects
    )

    if not subjects:
        await callback.message.edit_text(
            f"🎓 <b>{grade}-sinf</b>\n\n"
            "📚 Bu sinf uchun darslik hali "
            "yuklanmagan.\n\n"
            f"PDF fayllarni:\n"
            f"<code>books/{grade}-sinf/</code>\n"
            "papkasiga joylashtiring.\n\n"
            "🔄 /restart"
        )

        await callback.answer()
        return

    buttons = []

    for index, subject in enumerate(subjects):
        buttons.append(
            [
                InlineKeyboardButton(
                    text=f"📚 {subject}",
                    callback_data=f"subject:{index}"
                )
            ]
        )

    buttons.append(
        [
            InlineKeyboardButton(
                text="⬅️ Sinflar",
                callback_data="classes"
            )
        ]
    )

    await callback.message.edit_text(
        f"🎓 <b>{grade}-sinf</b>\n\n"
        "📚 Fanni tanlang:",
        reply_markup=InlineKeyboardMarkup(
            inline_keyboard=buttons
        )
    )

    await state.set_state(
        TeacherStates.choosing_subject
    )

    await callback.answer()


# ============================================================
# SUBJECT TANLASH
# ============================================================

@dp.callback_query(
    TeacherStates.choosing_subject,
    F.data.startswith("subject:")
)
async def subject_selected(
    callback: CallbackQuery,
    state: FSMContext
):
    data = await state.get_data()

    subjects = data.get("subjects", [])
    grade = data.get("grade")

    try:
        index = int(
            callback.data.split(":")[1]
        )

        subject = subjects[index]

    except (ValueError, IndexError):
        await callback.answer(
            "Fan topilmadi.",
            show_alert=True
        )
        return

    books = get_books(
        grade,
        subject
    )

    await state.update_data(
        subject=subject,
        books=[str(x) for x in books]
    )

    if not books:
        await callback.message.edit_text(
            f"📚 <b>{subject}</b>\n\n"
            "Bu fan uchun darslik PDF topilmadi.\n\n"
            f"PDFni:\n"
            f"<code>books/{grade}-sinf/{subject}/</code>\n"
            "papkasiga joylashtiring."
        )

        await callback.answer()
        return

    buttons = []

    for i, book in enumerate(books):
        buttons.append(
            [
                InlineKeyboardButton(
                    text=f"📖 {book.stem}",
                    callback_data=f"book:{i}"
                )
            ]
        )

    buttons.append(
        [
            InlineKeyboardButton(
                text="⬅️ Fanlar",
                callback_data="subjects"
            )
        ]
    )

    await callback.message.edit_text(
        f"🎓 {grade}-sinf\n"
        f"📚 <b>{subject}</b>\n\n"
        "📖 Darslikni tanlang:",
        reply_markup=InlineKeyboardMarkup(
            inline_keyboard=buttons
        )
    )

    await state.set_state(
        TeacherStates.choosing_book
    )

    await callback.answer()


# ============================================================
# BOOK TANLASH
# ============================================================

@dp.callback_query(
    TeacherStates.choosing_book,
    F.data.startswith("book:")
)
async def book_selected(
    callback: CallbackQuery,
    state: FSMContext
):
    data = await state.get_data()

    books = data.get("books", [])
    subject = data.get("subject")
    grade = data.get("grade")

    try:
        index = int(
            callback.data.split(":")[1]
        )

        book_path = Path(
            books[index]
        )

    except (ValueError, IndexError):
        await callback.answer(
            "Darslik topilmadi.",
            show_alert=True
        )
        return

    await state.update_data(
        book=str(book_path),
        book_name=book_path.stem
    )

    await callback.message.edit_text(
        f"🎓 <b>{grade}-sinf</b>\n"
        f"📚 <b>{subject}</b>\n"
        f"📖 <b>{book_path.stem}</b>\n\n"
        "📄 Qaysi sahifalardan test kerak?\n\n"
        "Masalan:\n"
        "<code>10-20</code>\n\n"
        "Yoki bitta sahifa:\n"
        "<code>15</code>"
    )

    await state.set_state(
        TeacherStates.entering_pages
    )

    await callback.answer()


# ============================================================
# PAGE INPUT
# ============================================================

@dp.message(
    TeacherStates.entering_pages
)
async def pages_entered(
    message: Message,
    state: FSMContext
):
    value = message.text.strip()

    parsed = parse_pages(value)

    if not parsed:
        await message.answer(
            "❌ Sahifa formati noto'g'ri.\n\n"
            "Masalan:\n"
            "<code>10-20</code>\n"
            "yoki\n"
            "<code>15</code>\n\n"
            "Maksimal oralig': 100 sahifa."
        )
        return

    start_page, end_page = parsed

    data = await state.get_data()

    book = Path(
        data.get("book", "")
    )

    grade = data.get("grade")
    subject = data.get("subject")
    book_name = data.get("book_name")

    if not book.exists():
        await message.answer(
            "❌ Darslik PDF topilmadi.\n"
            "Iltimos, /restart qiling."
        )
        return

    await message.answer(
        "📖 Darslik o'qilmoqda...\n"
        f"📄 Sahifalar: {start_page}-{end_page}\n\n"
        "⏳ Biroz kuting..."
    )

    source_text, total_pages, error = read_pdf(
        book,
        start_page,
        end_page
    )

    if error:
        await message.answer(
            f"❌ {error}"
        )
        return

    if not source_text.strip():
        await message.answer(
            "❌ Tanlangan sahifalardan matn topilmadi.\n\n"
            "Bu PDF skanerlangan rasm bo'lishi mumkin. "
            "Bunday PDF uchun OCR kerak bo'ladi.\n\n"
            "Boshqa matnli PDF bilan sinab ko'ring."
        )
        return

    if end_page > total_pages:
        await message.answer(
            f"ℹ️ Darslikda {total_pages} ta sahifa bor.\n"
            f"Siz so'ragan {end_page}-sahifa mavjud emas.\n"
            f"Shuning uchun {start_page}-{total_pages} "
            "oralig'i ishlatildi."
        )

    await message.answer(
        "🧠 <b>20 ta test tayyorlanmoqda...</b>\n\n"
        "📚 Testlar faqat tanlangan sahifalardagi "
        "matndan tuziladi."
    )

    user_id = message.from_user.id

    questions = generate_questions(
        source_text=source_text,
        count=20,
        user_id=user_id
    )

    if len(questions) < 20:
        await message.answer(
            f"⚠️ Bu sahifalardan faqat "
            f"<b>{len(questions)}</b> ta yangi test "
            "tuzish mumkin bo'ldi.\n\n"
            "20 ta test kerak bo'lsa, sahifa oralig'ini "
            "kengaytiring yoki boshqa sahifalarni tanlang."
        )

    if not questions:
        await message.answer(
            "❌ Test yaratish uchun yetarli matn topilmadi.\n\n"
            "Ko'proq sahifa tanlang."
        )
        return

    # Oxirgi ma'lumotlarni saqlash
    USER_DATA[user_id] = {
        "source_text": source_text,
        "questions": questions,
        "grade": grade,
        "subject": subject,
        "book": str(book),
        "book_name": book_name,
        "page_range": (
            f"{start_page}-{min(end_page, total_pages)}"
        ),
    }

    # State ichiga ham saqlaymiz
    await state.update_data(
        source_text=source_text,
        questions=questions,
        grade=grade,
        subject=subject,
        book=str(book),
        book_name=book_name,
        page_range=(
            f"{start_page}-{min(end_page, total_pages)}"
        )
    )

    await send_questions(
        message,
        questions
    )

    await message.answer(
        "🎉 <b>Testlar tayyor!</b>\n\n"
        "Quyidagi tugmalardan foydalanishingiz mumkin:",
        reply_markup=test_menu_keyboard()
    )


# ============================================================
# YANA 20 TA
# ============================================================

@dp.callback_query(
    F.data == "again"
)
async def again_questions(
    callback: CallbackQuery,
    state: FSMContext
):
    user_id = callback.from_user.id

    data = USER_DATA.get(user_id)

    if not data:
        data = await state.get_data()

    source_text = data.get("source_text")

    if not source_text:
        await callback.message.answer(
            "❌ Eski test ma'lumotlari topilmadi.\n"
            "Iltimos, /restart qiling."
        )

        await callback.answer()
        return

    await callback.message.answer(
        "🔄 <b>Yangi 20 ta test tayyorlanmoqda...</b>"
    )

    questions = generate_questions(
        source_text=source_text,
        count=20,
        user_id=user_id
    )

    if not questions:
        await callback.message.answer(
            "⚠️ Bu sahifalardan boshqa yangi savollar "
            "tuzish qiyin bo'ldi.\n\n"
            "📄 Boshqa sahifalarni tanlang."
        )

        await callback.answer()
        return

    data["questions"] = questions

    USER_DATA[user_id] = data

    await state.update_data(
        questions=questions
    )

    await send_questions(
        callback.message,
        questions
    )

    await callback.message.answer(
        "✅ Yangi testlar tayyor!",
        reply_markup=test_menu_keyboard()
    )

    await callback.answer()


# ============================================================
# WORD
# ============================================================

@dp.callback_query(
    F.data == "word"
)
async def word_callback(
    callback: CallbackQuery,
    state: FSMContext
):
    user_id = callback.from_user.id

    data = USER_DATA.get(user_id)

    if not data:
        data = await state.get_data()

    questions = data.get("questions")

    if not questions:
        await callback.answer(
            "Testlar topilmadi.",
            show_alert=True
        )
        return

    await callback.message.answer(
        "📄 <b>Word fayl tayyorlanmoqda...</b>"
    )

    try:
        filepath = create_word_file(
            user_id=user_id,
            questions=questions,
            grade=data.get("grade"),
            subject=data.get("subject"),
            book_name=data.get("book_name"),
            page_range=data.get("page_range")
        )

        document = FSInputFile(
            filepath
        )

        await callback.message.answer_document(
            document=document,
            caption=(
                "📄 <b>USTOZ AI TEST</b>\n\n"
                "Word fayl tayyor.\n"
                "Uni yuklab olib, o'zgartirishingiz "
                "yoki chop etishingiz mumkin."
            )
        )

        try:
            filepath.unlink()
        except Exception:
            pass

    except Exception as e:
        await callback.message.answer(
            f"❌ Word yaratishda xatolik:\n"
            f"<code>{e}</code>"
        )

    await callback.answer()


# ============================================================
# FANLAR
# ============================================================

@dp.callback_query(
    F.data == "subjects"
)
async def back_to_subjects(
    callback: CallbackQuery,
    state: FSMContext
):
    data = await state.get_data()

    grade = data.get("grade")

    if not grade:
        await callback.answer()
        return

    subjects = get_subjects(grade)

    await state.update_data(
        subjects=subjects
    )

    if not subjects:
        await callback.message.edit_text(
            f"🎓 {grade}-sinf\n\n"
            "📚 Darsliklar topilmadi."
        )

        await callback.answer()
        return

    buttons = []

    for index, subject in enumerate(subjects):
        buttons.append(
            [
                InlineKeyboardButton(
                    text=f"📚 {subject}",
                    callback_data=f"subject:{index}"
                )
            ]
        )

    buttons.append(
        [
            InlineKeyboardButton(
                text="⬅️ Sinflar",
                callback_data="classes"
            )
        ]
    )

    await callback.message.edit_text(
        f"🎓 <b>{grade}-sinf</b>\n\n"
        "📚 Fanni tanlang:",
        reply_markup=InlineKeyboardMarkup(
            inline_keyboard=buttons
        )
    )

    await state.set_state(
        TeacherStates.choosing_subject
    )

    await callback.answer()


# ============================================================
# SINFLAR
# ============================================================

@dp.callback_query(
    F.data == "classes"
)
async def back_to_classes(
    callback: CallbackQuery,
    state: FSMContext
):
    await state.clear()

    await callback.message.edit_text(
        "🎓 <b>Sinfni tanlang:</b>",
        reply_markup=class_keyboard()
    )

    await state.set_state(
        TeacherStates.choosing_class
    )

    await callback.answer()


# ============================================================
# HOME
# ============================================================

@dp.callback_query(
    F.data == "home"
)
async def home_callback(
    callback: CallbackQuery,
    state: FSMContext
):
    await state.clear()

    await callback.message.edit_text(
        "🏠 <b>USTOZ AI</b>\n\n"
        "O'qituvchi uchun test generator.\n\n"
        "🎓 Sinfni tanlang:",
        reply_markup=class_keyboard()
    )

    await state.set_state(
        TeacherStates.choosing_class
    )

    await callback.answer()


# ============================================================
# UNKNOWN MESSAGE
# ============================================================

@dp.message()
async def unknown_message(message: Message):
    await message.answer(
        "🤖 Men sizni tushunmadim.\n\n"
        "Boshlash uchun /start ni bosing.\n"
        "Yordam: /help\n"
        "Qayta boshlash: /restart"
    )


# ============================================================
# MAIN
# ============================================================

async def main():
    print("=" * 45)
    print("        USTOZ AI ISHGA TUSHDI")
    print("=" * 45)

    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())